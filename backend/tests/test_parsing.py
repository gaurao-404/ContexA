from io import BytesIO
from docx import Document
from app.parsing import parse_document


def test_parse_plain_text():
    pages, text = parse_document('sample.txt', b'ContexA review workflow')
    assert pages[0]['page'] == 1
    assert 'review workflow' in text


def test_parse_docx():
    document = Document()
    document.add_paragraph('Use approved claims before publication.')
    output = BytesIO(); document.save(output)
    pages, text = parse_document('sample.docx', output.getvalue())
    assert 'approved claims' in text
    assert pages[0]['text'] == text


def test_parse_pptx_reports_picture_without_alt_text():
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image

    image_bytes = BytesIO()
    Image.new('RGB', (32, 32), color=(120, 80, 160)).save(image_bytes, format='PNG')
    image_bytes.seek(0)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(image_bytes, Inches(1), Inches(1), width=Inches(1), height=Inches(1))
    output = BytesIO(); presentation.save(output)
    pages, _ = parse_document('visual-check.pptx', output.getvalue())
    assert len(pages[0]['missing_alt_text']) == 1
    assert pages[0]['missing_alt_text'][0]['name']
