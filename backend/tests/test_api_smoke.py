from __future__ import annotations

import fitz
import pytest
from fastapi.testclient import TestClient

from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as value:
        yield value


def login(client: TestClient, email: str, password: str) -> dict[str, str]:
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_upload_review_gate_revision_export_and_audit(client: TestClient):
    admin = login(client, "admin@contexa.local", "ContexA!2026")
    creator = login(client, "creator@contexa.local", "Creator!2026")
    upload = client.post("/api/documents/upload", headers=admin, files={
        "file": ("proof-claim.txt", b"Contexa says we are the #1 fastest platform to boost revenue by 400%.", "text/plain")
    })
    assert upload.status_code == 201, upload.text
    doc_id = upload.json()["id"]
    detail = client.get(f"/api/documents/{doc_id}", headers=admin)
    assert detail.status_code == 200
    findings = detail.json()["findings"]
    assert {f["rule_code"] for f in findings} >= {"CLM-204", "LGL-002"}
    critical = [f for f in findings if f["severity"] == "critical"]
    assert len(critical) >= 1

    forbidden = client.post(f"/api/findings/{critical[0]['id']}/action", headers=creator,
        json={"status": "accepted"})
    assert forbidden.status_code == 403

    # An attempted edit without exact replacement copy is rejected.
    warning = next(f for f in findings if f["rule_code"] == "TERM-008")
    invalid_edit = client.post(f"/api/findings/{warning['id']}/action", headers=admin,
        json={"status": "edited"})
    assert invalid_edit.status_code == 422
    revised = client.post(f"/api/findings/{warning['id']}/action", headers=admin,
        json={"status": "edited", "note": "Reviewed revised copy", "applied_text": "ContexA is designed to help teams review approved brand content."})
    assert revised.status_code == 200, revised.text
    assert revised.json()["finding"]["applied_text"].startswith("ContexA is designed")

    for finding in critical:
        res = client.post(f"/api/findings/{finding['id']}/action", headers=admin,
            json={"status": "dismissed", "note": "Legal approved this statement for the internal demo."})
        assert res.status_code == 200, res.text
    blocked = client.post(f"/api/documents/{doc_id}/publish", headers=admin, json={"override_note": ""})
    assert blocked.status_code == 422
    approved = client.post(f"/api/documents/{doc_id}/publish", headers=admin,
        json={"override_note": "Legal reviewer accepted the documented exceptions for demo purposes."})
    assert approved.status_code == 200, approved.text
    assert approved.json()["document"]["status"] == "PUBLISH_READY"

    export = client.get(f"/api/documents/{doc_id}/export?fmt=csv", headers=admin)
    assert export.status_code == 200 and "applied_text" in export.text
    audit = client.get("/api/audit?limit=30", headers=admin)
    assert audit.status_code == 200
    assert any(item["action"] == "document.publish_ready" for item in audit.json())


def test_pdf_preview_is_real_rendered_content(client: TestClient):
    admin = login(client, "admin@contexa.local", "ContexA!2026")
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    page.insert_text((55, 90), "The #1 fastest platform to boost revenue by 400%.", fontsize=16, color=(0.78, 0.78, 0.78))
    payload = pdf.tobytes()
    pdf.close()
    upload = client.post("/api/documents/upload", headers=admin,
        files={"file": ("sample.pdf", payload, "application/pdf")})
    assert upload.status_code == 201, upload.text
    doc_id = upload.json()["id"]
    preview = client.get(f"/api/documents/{doc_id}/preview/1", headers=admin)
    assert preview.status_code == 200, preview.text[:200]
    assert preview.headers["content-type"].startswith("image/png")
    assert preview.content.startswith(b"\x89PNG")


def test_write_time_copy_check_uses_approved_workspace_rules(client: TestClient):
    admin = login(client, "admin@contexa.local", "ContexA!2026")
    response = client.post("/api/check-copy", headers=admin, json={
        "text": "Contexa is the #1 fastest platform to boost revenue by 400%. Our revolutionary plan is only $19/month — save 30% today!"
    })
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["mode"] == "deterministic"
    assert payload["findings_count"] >= 4
    codes = {finding["rule_code"] for finding in payload["findings"]}
    assert {"CLM-204", "LGL-002", "TERM-001", "TERM-008", "PRC-002"}.issubset(codes)
    assert "cannot determine freshness" in next(f["suggested_fix"] for f in payload["findings"] if f["rule_code"] == "PRC-002")


def test_pptx_upload_flags_picture_with_filename_only_alt_text(client: TestClient):
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image

    admin = login(client, "admin@contexa.local", "ContexA!2026")
    image = BytesIO(); Image.new("RGB", (32, 32), color=(117, 88, 238)).save(image, format="PNG"); image.seek(0)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(image, Inches(1), Inches(1), width=Inches(1), height=Inches(1))
    output = BytesIO(); presentation.save(output)
    uploaded = client.post("/api/documents/upload", headers=admin,
        files={"file": ("campaign-slide.pptx", output.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert uploaded.status_code == 201, uploaded.text
    detail = client.get(f"/api/documents/{uploaded.json()['id']}", headers=admin)
    assert detail.status_code == 200
    accessibility = [f for f in detail.json()["findings"] if f.get("rule_code") == "A11Y-001"]
    assert accessibility, "Expected a verified missing/meaningless alt-text finding for the PowerPoint picture."


def test_guideline_import_creates_paused_reviewable_rules(client: TestClient):
    admin = login(client, "admin@contexa.local", "ContexA!2026")
    creator = login(client, "creator@contexa.local", "Creator!2026")
    denied = client.post("/api/guidelines/import", headers=creator,
        files={"file": ("guide.txt", b"Always support quantified claims with evidence before publication.", "text/plain")})
    assert denied.status_code == 403

    source = (b"Never use world-class in marketing copy.\n"
              b"Pricing must match the current approved price book and campaign validity window.\n"
              b"Claims about performance or revenue require evidence and legal review before publication.\n")
    response = client.post("/api/guidelines/import", headers=admin,
        files={"file": ("brand-policy.txt", source, "text/plain")})
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["created_rules"] == 3
    assert all(not rule["is_active"] for rule in payload["rules"])
    rules = client.get("/api/rules?q=brand-policy.txt", headers=admin)
    imported = [rule for rule in rules.json() if rule["id"] in {candidate["id"] for candidate in payload["rules"]}]
    assert len(imported) == 3
    assert all(rule["is_active"] is False for rule in imported)
    assert any(rule["pattern"] for rule in imported), "Simple explicit prohibitions should yield a reviewable deterministic pattern candidate."


def test_demo_seed_is_repeatable_and_does_not_require_local_ai(client: TestClient):
    admin = login(client, "admin@contexa.local", "ContexA!2026")
    creator = login(client, "creator@contexa.local", "Creator!2026")
    denied = client.post("/api/demo/seed", headers=creator)
    assert denied.status_code == 403

    first = client.post("/api/demo/seed", headers=admin)
    assert first.status_code == 200, first.text
    payload = first.json()
    assert payload["ai_mode"] == "deterministic_only"
    assert len(payload["documents"]) == 3
    assert "No Ollama request" in payload["notice"]
    assert {doc["filename"] for doc in payload["documents"]} >= {
        "DEMO_Q4_Launch_Campaign.txt", "DEMO_Product_Naming_Review.txt", "DEMO_Pricing_Refresh.txt"
    }

    second = client.post("/api/demo/seed", headers=admin)
    assert second.status_code == 200, second.text
    assert second.json()["created_documents"] == 0
    assert len(second.json()["documents"]) == 3
